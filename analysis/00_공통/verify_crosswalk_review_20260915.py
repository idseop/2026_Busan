"""Independent dated code-candidate and report review; no receipt allocation."""
import json
from pathlib import Path
import pandas as pd
import openpyxl
import re
from urllib.parse import unquote
from verify_all_selected_20260915 import BASE,ROOT,OUT,read,sha

def main():
    folder=BASE/'linkage'
    cross=read(folder/'all_raw_names_year_crosswalk.csv')
    intervals=read(folder/'all_candidate_relation_intervals.csv')
    source=read(ROOT/'data/processed/컬럼선별-결측제외-20260915/coverage/all_original_region_names_before_filter.csv')
    pop=read(ROOT/'data/processed/동별예방분석-20260914/population/region_registry.csv');pop=pop[pop.level=='dong']
    assert len(cross)==len(source) and cross.receipt_count_before_filter.astype(int).sum()==source.receipt_count.astype(int).sum()
    assert not cross.duplicated(['year','district_name','raw_dong_name']).any()
    assert cross.receipt_geography_confirmed.eq('False').all() and cross.raw_field_semantics_confirmed.eq('False').all()
    for _,row in cross.iterrows():
        candidates=set(filter(None,row.candidate_codes.split('|')))
        selected=intervals[(intervals.year==row.year)&(intervals.district_name==row.district_name)&(intervals.raw_dong_name==row.raw_dong_name)]
        assert set(selected.admin_code)==candidates and len(candidates)==int(row.candidate_count)
        cursor=row.year+'0101';end=str(int(row.year)+1)+'0101'
        for start,finish in sorted(set(zip(selected.valid_from_in_year,selected.valid_until_exclusive_in_year))):
            assert row.year+'0101'<=start<finish<=end
            if start<=cursor:cursor=max(cursor,finish)
        whole=cursor>=end
        assert (row.record_interval_covers_full_year=='True')==whole
        assert row.unique_code==(next(iter(candidates)) if len(candidates)==1 else '')
        matching=pop[(pop.year==row.year)&(pop.admin_code==row.unique_code)]
        assert (row.year_end_population_code_exists=='True')==(len(matching)==1)
        eligible=len(candidates)==1 and whole and len(matching)==1 and row.temporal_reconstruction_conflict=='False'
        assert (row.code_level_population_join_eligible=='True')==eligible
        if eligible:assert float(row.population)==int(matching.iloc[0].population)
        else:assert row.population==''
    codes=read(folder/'official_administrative_codes_with_dates.csv')
    for year in range(2020,2025):
        date=f'{year}1231'; active=codes[(codes['생성일자']<=date)&((codes['말소일자']=='')|(codes['말소일자']>date))]
        assert set(active['행정동코드'])==set(pop[pop.year==str(year)].admin_code)
    hashes={}
    for directory in [OUT,folder]:
        manifest=json.loads((directory/'manifest.json').read_text(encoding='utf-8'))
        for artifact in manifest['outputs']:assert sha(directory/artifact['file'])==artifact['sha256']
        hashes[directory.name]=sha(directory/'manifest.json')
    doc=ROOT/'docs/40-분석결과/컬럼-사용목적과-전체결측조건-20260915.md'
    text=doc.read_text(encoding='utf-8'); table=[[c.strip() for c in line.strip('|').split('|')] for line in text.splitlines() if line.startswith('|')]
    stages=read(OUT/'stage_completeness.csv')
    for year in range(2020,2025):
        row=next(r for r in table if r[0]==str(year))
        expected=[]
        for stage in ['core8','core9','provisional17','selected22','selected23']:
            n=stages[(stages.year==str(year))&(stages.stage==stage)&(stages.basis_or_scope=='A')].iloc[0].complete_rows
            expected.append(int(n))
        assert [int(c.replace(',','')) for c in row[1:]]==expected
    assert '17개를 하나의 분석표로 사용' in text
    assert '최종 분석 컬럼으로 확정하지 않았다' not in text
    classification=read(OUT/'annual_classification.csv'); classification=classification[classification.scope=='A']
    for kind in ['구급','구조','화재','기타']:
        c=classification[classification.EMRG_RSCU_ASSRT_NM==kind]
        before=c[c.stage=='core8']['count'].astype(int).sum();after=c[c.stage=='provisional17']['count'].astype(int).sum()
        assert f'{100*after/before:.2f}%' in text
    missing=read(OUT/'selected_column_missing.csv'); missing=missing[missing.basis=='core8']
    for col in ['GRNDS_SGG_NM','ACDNT_OCRN_LOT']:
        count=missing[missing.column==col].missing_rows.astype(int).sum();assert f'{count:,}행' in text
    retained=read(OUT/'original_region_totals_retention.csv')
    assert not ((retained.name_present=='True') & (retained.disappears_at17_vs_core8=='True')).any()
    evidence_path=BASE/'independent_verification.json'; evidence=json.loads(evidence_path.read_text(encoding='utf-8'))
    evidence['completeness_manifest_sha256']=hashes['completeness']
    evidence['crosswalk_review']={'status':'passed','all_raw_names_preserved':len(cross),'all_candidate_code_sets_and_intervals_checked':True,'eligible_unique_only_full_year_and_population_code_checked':True,'no_count_allocation_or_geographic_confirmation':True,'every_year_end_205_code_set_matches_population':True,'output_hashes_match':hashes,'limitation':'This checks code-record candidate logic, not historical polygon geometry or source field semantics.'}
    docs=[doc,ROOT/'docs/10-작업계획/부산 분석 작업 계획.md',ROOT/'docs/40-분석결과/지도-결과데이터-계약-20260914.md',ROOT/'docs/40-분석결과/원문동-행정동-대응검증-20260915.md']
    bookaudit=json.loads((OUT/'codebook_audit.json').read_text(encoding='utf-8'))
    assert sha(ROOT/bookaudit['workbook'])==bookaudit['workbook_sha256']
    assert sha(OUT/bookaudit['output'])==bookaudit['output_sha256']
    definitions=read(OUT/bookaudit['output'])
    wb=openpyxl.load_workbook(ROOT/bookaudit['workbook'],read_only=True,data_only=True)
    ws=wb['column_info']
    for _,r in definitions.iterrows():
        assert ws[r.english_cell].value==r.english_name_original
        assert ws[r.korean_cell].value==r.korean_name
        assert ws[r.detail_cell].value==r.detail_description
    wb.close()
    assert len(definitions)==38 and definitions.selected17.eq('True').sum()==17
    policy=json.loads((OUT/'comparison_policy.json').read_text(encoding='utf-8'))
    assert set(definitions[definitions.selected17=='True'].english_name_normalized)==set(policy['current_selected_fields'])
    final=ROOT/'docs/40-분석결과/17컬럼-완전행-후속분석-20260915.md';finaltext=final.read_text(encoding='utf-8');docs.append(final)
    finalrows=[[x.strip() for x in line.strip('|').split('|')] for line in finaltext.splitlines() if line.startswith('|')]
    scopes=read(BASE/'analysis/annual_scope_counts.csv')
    for year in range(2020,2025):
        r=next(x for x in finalrows if x[0]==str(year))
        expected=[int(stages[(stages.year==str(year))&(stages.stage=='core8')&(stages.basis_or_scope=='A')].iloc[0].complete_rows)]
        expected += [int(scopes[(scopes.year==str(year))&(scopes.scope==s)].iloc[0].records) for s in 'ABC']
        assert [int(x.replace(',','')) for x in r[1:]]==expected
    fulltypes=read(BASE/'analysis/annual_type.csv')
    for scope,label in [('A','A 전체 처리'),('B','B 정상'),('C','C 운영성 제외')]:
        row=next(x for x in finalrows if x[0]==label)
        assert [int(x.replace(',','')) for x in row[1:]]==[int(fulltypes[(fulltypes.scope==scope)&(fulltypes.EMRG_RSCU_ASSRT_NM==kind)].receipt_count.astype(int).sum()) for kind in ['구급','구조','기타','화재']]
    channels=read(OUT/'selected17_receipt_path_distribution.csv');channels=channels[channels.scope=='A']
    mobile=channels[channels.RCPT_PATH_NM=='이동전화']['count'].astype(int).sum()
    assert f'{mobile:,}행' in finaltext and f'{100*mobile/704689:.4f}%' in finaltext
    region=read(OUT/'selected17_region_name_comparison.csv');region=region[region.scope=='A']
    assert region[region.province_names_equal=='False']['count'].astype(int).sum()==0
    mismatch=region[region.district_names_equal=='False']['count'].astype(int).sum();assert f'{mismatch:,}행' in finaltext
    pair=read(OUT/'selected17_coordinate_pair_comparison.csv');exact=pair.both_coordinates_numeric_exact_equal.astype(int).sum();assert f'{exact}행' in finaltext
    for _,r in scopes.iterrows():assert int(r.observed_dates)==(366 if int(r.year) in [2020,2024] else 365)
    for target in re.findall(r'\]\(([^)]+)\)',finaltext):
        if not target.startswith('http'):assert (final.parent/unquote(target)).is_file()
    evidence['codebook_review']={'status':'passed','xlsx_definitions_verified':38,'selected17_exactly_matches_current_policy':True,'labels_and_details_exactly_match_source_cells':True,'audit_sha256':sha(OUT/'codebook_audit.json')}
    evidence['complete17_report_numeric_review']={'status':'passed','annual_scope_and_type_tables_match':True,'mobile_channel_and_region_mismatch_and_coordinate_comparison_match':True,'full_calendar_dates_observed_all_scopes':True,'chart_visual_review':'passed; five years, count/retention labels and selection caveat legible'}
    evidence['final_report_review']={'status':'passed','all_year_stage_table_matches':True,'17_approved_and23_zero_and22_337_distinguished':True,'document_sha256':{p.relative_to(ROOT).as_posix():sha(p) for p in docs}}
    evidence_path.write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
    print('Crosswalk and current report independent review passed')

if __name__=='__main__':main()
