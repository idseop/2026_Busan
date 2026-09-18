"""Independent readback: no imports from the analysis/writer implementation."""
from pathlib import Path
import json, hashlib, re
import openpyxl, pandas as pd, numpy as np
import pymupdf

ROOT=Path(__file__).resolve().parents[2];O=ROOT/'data/processed/신고보건심화-20260917';V=O/'verification';V.mkdir(exist_ok=True)
checks=[]
def check(name,condition,detail=None):
    checks.append(dict(name=name,passed=bool(condition),detail=detail))
def hashfile(p):return hashlib.sha256(p.read_bytes()).hexdigest()
spec=json.loads((O/'analysis-specification.json').read_text(encoding='utf-8'))
for rec in spec['inputs']:check('input:'+rec['file'],hashfile(ROOT/rec['file'])==rec['sha256'])
for manifest in ['sources/chs-download-manifest.json','sources/chs2023-download-manifest.json','sources/service-source-manifest.json']:
    for rec in json.loads((O/manifest).read_text(encoding='utf-8')):
        if rec.get('sha256'):check('source:'+rec.get('district',rec.get('id',''))+':'+manifest,hashfile(ROOT/rec['file'])==rec['sha256'])
        for item in rec.get('extracted',[]):check('extracted:'+item['file'],hashfile(ROOT/item['file'])==item['sha256'])

s=pd.read_csv(O/'health-district-summary-2024.csv')
long=pd.read_csv(O/'health-indicators-all-published-groups.csv')
t=pd.read_csv(O/'health-same-indicators-2023-2024.csv')
check('rows',len(long)==6980 and len(s)==160 and len(t)==30)
check('32 indicators across 5 districts',s.groupby('district').size().eq(32).all() and s.district.nunique()==5)
check('unique outputs',not long.duplicated(['district','year','indicator','group','sex']).any() and not t.duplicated(['district','year','indicator']).any())
check('suppression preserved',long[long.published_value=='-(-)'].estimate.isna().all() and len(long[long.published_value=='-(-)'])>0)
check('survey values valid',long.dropna(subset=['estimate']).estimate.ge(0).all() and long[long.unit=='%'].estimate.dropna().le(100).all())
check('CI consistent',np.allclose(long.dropna(subset=['estimate']).lower_approx95,np.maximum(0,long.dropna(subset=['estimate']).estimate-1.96*long.dropna(subset=['estimate']).se)))
check('rate precision flags',long[(long.rse>=25)&long.estimate.notna()].status.eq('정밀도 주의').all())

# The author reads 2024 hidden data rows. Verify key published values using visible
# tables selected by titles and row labels, including the necessary-care denominator.
books={}
for gu in s.district.unique():
    path=ROOT/s[s.district==gu].iloc[0].source_file;w=openpyxl.load_workbook(path,read_only=True,data_only=True);books[(gu,2024)]=w
    for indicator in ['연간 미충족의료율(병의원)','고혈압 진단 경험자의 치료율_30','고혈압 관리교육 이수율_30','당뇨병 진단 경험자의 치료율_30','당뇨병 관리교육 이수율_30','연간 사고 및 중독 경험률']:
        name='<표>'+indicator.replace('_30','').replace(' ','')
        sheet=next(ws for ws in w if ws.title.replace(' ','')==name)
        found=[]
        for cells in sheet:
            # visible label resides in column B in the 2024 workbook
            label=str(cells[1].value or '')
            if (indicator.endswith('_30') and label.startswith('30세 이상')) or (not indicator.endswith('_30') and label=='전체'):
                n,rate,rse=[c.value for c in cells[2:5]]
                m=re.fullmatch(r'(\d+(?:\.\d+)?)\((\d+(?:\.\d+)?)\)',str(rate))
                if m:found.append((int(n),float(m[1]),float(m[2]),float(rse)))
        q=s[(s.district==gu)&(s.indicator==indicator)].iloc[0]
        check('visible2024:'+gu+indicator,len(found)==1 and np.allclose(found[0],[q.n,q.estimate,q.se,q.rse]))
    # Validate every all-sex summary row against its original published data cell.
    ws=w['data'];headers={c.value:c.column for c in ws[1]}
    for row in s[s.district==gu].itertuples():
        ncol,vcol,rcol=('n0','m_se0','rse0') if ws.cell(row.source_excel_row,headers['type']).value=='type1' else ('N','M_SE','rse')
        check('all summary readback:'+gu+row.indicator,str(ws.cell(row.source_excel_row,headers[vcol]).value)==row.published_value and int(ws.cell(row.source_excel_row,headers[ncol]).value)==row.n)
    reason=s[(s.district==gu)&s.indicator.str.startswith('미충족의료 이유(병의원)_')]
    unmet=s[(s.district==gu)&s.indicator.eq('연간 미충족의료율(병의원)')].iloc[0]
    check('reason denominator:'+gu,reason.n.eq(unmet.n).all())
    check('reason rounding subtotal:'+gu,abs(reason.estimate.sum()-unmet.estimate)<=.5)

# 2023 table authoring is verified against independent PDF narrative where available,
# plus the separate hidden-data layout in the Busanjin original workbook.
for gu in ['북구','사하구']:
    doc=pymupdf.open(O/f'sources/chs2023-{gu}.pdf')
    text='\n'.join(p.get_text() for p in doc)
    pattern=r'연간\s*미충족의료율\(병의원\)은\s*([0-9.]+)%'
    found=re.findall(pattern,text)
    expected=t[(t.district==gu)&t.year.eq(2023)&t.indicator.eq('연간 미충족의료율(병의원)')].iloc[0].estimate
    check('independent PDF2023:'+gu,len(found)>0 and float(found[0])==expected)
old=pd.read_excel(O/'sources/chs2023-부산진구.xlsx',sheet_name=0,dtype=str)
for row in t[(t.district=='부산진구')&t.year.eq(2023)].itertuples():
    q=old[(old.kor_var==row.indicator)&(old.class_nm=='전체')].iloc[0]
    check('independent hidden2023:'+row.indicator,q.m_se0==row.published_value and int(q.n0)==row.n)
for gu in ['북구','사하구','부산진구']:
    for year in [2023,2024]:
        for disease in ['고혈압','당뇨병']:
            q=t[(t.district==gu)&t.year.eq(year)&t.indicator.str.startswith(disease)]
            check(f'education treatment shared denominator:{gu}{year}{disease}',q.n.nunique()==1 and len(q)==2)

base=ROOT/'data/processed/신고주민연결심화-20260917/direction-connections'
expected=pd.read_csv(base/'major-burden-top10.csv');link=pd.read_csv(O/'selected-region-evidence-links.csv')
for r in link.itertuples():
    x=expected[(expected.district==r.district)&(expected.rawDong==r.rawDong)&(expected.subtype==r.subtype)].iloc[0]
    check('unchanged call selection:'+r.rawDong+r.subtype,all(getattr(r,c)==x[c] for c in ['countP','countRankWithinSubtype','aboveRestBusan30','aboveRestDistrict30','2020','2021','2022','2023','2024'] if not c.isnumeric()))
    check('5 years sum:'+r.rawDong+r.subtype,sum(x[str(y)] for y in range(2020,2025))==r.countP)
    check('health attribution scope:'+r.rawDong,'구 전체 조사' in r.healthSpatialScope)
pop=pd.read_csv(base/'major-burden-population-candidates.csv');pop=pop[(pop.year==2024)&pop.rawDong.isin(['다대동','금곡동','부전동'])]
check('all resident ages kept',len(pop)==5 and np.allclose(pop[[f'age_{i}' for i in range(101)]].sum(axis=1),pop.residentTotal))
check('no incident geography upgraded',not pop.geographyConfirmed.any())
services=json.loads((O/'sources/service-source-manifest.json').read_text(encoding='utf-8'))
check('gated EMS explicitly unacquired',all(r.get('dataDownloaded') is False and r.get('loginGuardPresent') is True for r in services if r['id'].startswith('119')))
proposal=pd.read_csv(O/'regional-improvement-and-evaluation.csv')
check('conditional proposals',len(proposal)==3 and proposal.precondition.notna().all() and proposal.effect.str.contains('미검증|미확정|확정하지|없음').all())
check('separate prevention scope',not proposal.to_csv(index=False).count('벌집'))
for p in O.glob('*.csv'):
    header=pd.read_csv(p,nrows=0).columns
    check('aggregate only:'+p.name,not any(c in header for c in ['DCLR_RCPT_NO','ACDNT_OCRN_LOT','ACDNT_OCRN_LAT','ptn_no','rptp_no']))
for p in [ROOT/'docs/40-분석결과/부산-119-신고인구에서-보건이용과-보완으로-20260917.md',ROOT/'output/부산119-신고보건심화-20260917.md']:
    content=p.read_text(encoding='utf-8')
    for target in re.findall(r'\]\(([^)]+)\)',content):
        if target.startswith('http'):continue
        check('report link:'+target,(p.parent/target).resolve().exists())
for name in ['01-의료이용-추가근거','02-기존치료와-관리교육','03-두해-보건지표-검토']:
    for ext in ['png','svg']:
        p=ROOT/'figures/신고보건심화-20260917'/f'{name}.{ext}'
        check('figure:'+p.name,p.exists() and p.stat().st_size>10000)
result=dict(passed=all(c['passed'] for c in checks),checks=len(checks),results=checks,
    scope='별도 실행; 작성 모듈 미호출. 주요2024 지표는 공개 표시표와 재대조,2023 의료 미이용은 독립 PDF 서술 대조. 시범 시행·정책 효과 검증 아님')
(V/'verification.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
(V/'verification.md').write_text('# 신고보건심화 검증\n\n'+f"전체 {len(checks)}항목: {'PASS' if result['passed'] else 'FAIL'}\n\n"+result['scope']+'\n\n'+'\n'.join(f"- {'PASS' if c['passed'] else 'FAIL'}: {c['name']}" for c in checks),encoding='utf-8')
print(json.dumps(dict(passed=result['passed'],checks=len(checks),failed=[c for c in checks if not c['passed']]),ensure_ascii=False))
assert result['passed']
